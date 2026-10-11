"""`sz import <folder>`: bring in everything you produced before Setz existed.

Point it at a folder of exports (e.g. downloaded from your Claude projects "Settleezy", "Partnership and outreach"
and "Social media", or any spreadsheets/docs). It understands:

  CSV / XLSX / JSON   contacts, leads, partner lists  -> leads table (+ partners when the status says signed/live)
  MD / TXT / DOCX     outreach emails & templates     -> library/outreach (drafts reuse what worked)
                      social posts, captions, scripts -> library/social   (content ideas, growth review)
                      strategy notes, reports         -> library/notes    (growth review context)

Running it again is safe: leads are merged by name, documents are de-duplicated by content.
"""

from __future__ import annotations

import csv
import hashlib
import json
import re
from pathlib import Path
from typing import Any, Iterable

from .config import Config
from .db import DB

COLUMNS = {
    "name": ["name", "company", "business", "venue", "organisation", "organization", "brand", "partner", "university", "firma",
             "unternehmen", "lead", "account", "store", "restaurant", "cafe", "café"],
    "email": ["email", "e-mail", "mail", "email address", "e-mail-adresse"],
    "phone": ["phone", "telefon", "tel", "mobile", "handy", "phone number"],
    "instagram": ["instagram", "ig", "insta", "handle", "instagram handle"],
    "website": ["website", "url", "web", "site", "webseite", "homepage", "link"],
    "category": ["category", "industry", "branche", "kategorie", "segment", "sector"],
    "kind": ["kind", "type", "lead type", "partner type", "typ"],
    "status": ["status", "stage", "pipeline", "phase"],
    "notes": ["notes", "note", "comment", "comments", "notizen", "remarks", "next step"],
    "contact_name": ["contact", "contact person", "contact name", "ansprechpartner", "person", "first name", "full name"],
    "city": ["city", "stadt", "location", "ort"],
}
STATUS_MAP = [   # first match wins: negations before anything that could match "interested" / "partner"
    (r"\b(not interested|no interest|uninterested|declined|rejected|lost|kein interesse|nicht interessiert|abgesagt|absage|"
     r"abgelehnt|no)\b", "lost", None),
    (r"^\s*(new|open|to ?do|not (yet )?contacted|nicht kontaktiert|noch nicht|neu|offen)\b", None, None),
    (r"\b(live|active|aktiv|onboarded|launched)\b", "partner", "live"),
    (r"\b(signed|contract signed|agreed|won|partner|zugesagt|unterschrieben)\b", "partner", "agreed"),   # "contract sent" is not signed
    (r"meeting|call booked|termin", "meeting", None),
    (r"replied|responded|interested|antwort|interessiert", "replied", None),
    (r"contacted|sent|emailed|reached|kontaktiert|angeschrieben|follow", "contacted", None),
]
STATUS_ORDER = ["new", "drafted", "contacted", "replied", "meeting", "partner"]
KIND_HINTS = [
    (r"uni|hochschule|college|school|international office|akademie", "university"),
    (r"housing|wohn|apartment|coliving|residence|room", "housing"),
    (r"brand|unidays|student ?beans|marke", "brand"),
    (r"bank|insurance|versicherung|sim|telecom|mobile|service", "service"),
]


def _norm_header(h: str) -> str:
    return re.sub(r"[^a-z0-9äöüß -]", "", (h or "").strip().lower())


def _map_columns(headers: Iterable[str]) -> dict[str, str]:
    """Header -> field: exact names first ("Partner type" is the kind, not the name), then any e-mail column
    ("Contact email"), then prefixes ("Company name")."""
    hs = [(h, _norm_header(h)) for h in headers]
    out: dict[str, str] = {}

    def free(h: str) -> bool:
        return h not in out.values()

    for h, nh in hs:
        for field, names in COLUMNS.items():
            if field not in out and nh in names:
                out[field] = h
                break
    for h, nh in hs:
        if "email" not in out and free(h) and "mail" in nh:
            out["email"] = h
    for h, nh in hs:
        if not free(h):
            continue
        for field, names in COLUMNS.items():
            if field not in out and any(nh.startswith(n + " ") for n in names):
                out[field] = h
                break
    return out


def _rows_from_file(path: Path) -> list[dict[str, Any]]:
    suf = path.suffix.lower()
    if suf == ".csv":
        text = path.read_text(encoding="utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t") if text.strip() else csv.excel
        except csv.Error:   # a single column (no delimiter to find)
            dialect = csv.excel
        return list(csv.DictReader(text.splitlines(), dialect=dialect))
    if suf == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(data, dict):
            data = next((v for v in data.values() if isinstance(v, list)), [data])
        return [r for r in data if isinstance(r, dict)]
    if suf == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("Install openpyxl to import .xlsx (pip install openpyxl), or save the sheet as CSV.") from exc
        rows: list[dict[str, Any]] = []
        wb = load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            it = ws.iter_rows(values_only=True)
            header = next(it, None)
            if not header:
                continue
            for r in it:
                if any(c not in (None, "") for c in r):
                    rows.append({str(h or f"col{i}"): ("" if c is None else str(c)) for i, (h, c) in enumerate(zip(header, r))})
        return rows
    return []


def _kind_for(row: dict, mapping: dict, hint: str) -> str:
    raw = (row.get(mapping.get("kind", ""), "") or "") + " " + (row.get(mapping.get("category", ""), "") or "") + " " + hint
    for rx, kind in KIND_HINTS:
        if re.search(rx, raw, re.I):
            return kind
    return "merchant"


def import_table(db: DB, path: Path) -> dict[str, int]:
    from . import leads as leads_mod
    from .ops import ensure_partner_from_lead, update_partner

    rows = _rows_from_file(path)
    if not rows:
        return {"rows": 0, "leads_new": 0, "leads_updated": 0, "partners": 0}
    mapping = _map_columns(rows[0].keys())
    if "name" not in mapping and "email" not in mapping:
        return {"rows": len(rows), "skipped": len(rows), "leads_new": 0, "leads_updated": 0, "partners": 0}
    hint = path.stem.lower()
    stats = {"rows": len(rows), "leads_new": 0, "leads_updated": 0, "partners": 0}
    for row in rows:
        get = lambda f: str(row.get(mapping.get(f, ""), "") or "").strip()  # noqa: E731
        name = get("name") or _name_from_email(get("email"), get("contact_name"))
        if not name:
            continue
        kind = _kind_for(row, mapping, hint)
        ig = get("instagram")
        if ig and not ig.startswith("@") and "instagram.com" not in ig:
            ig = "@" + ig
        lid, created = leads_mod.upsert(db, name, kind, source=f"import:{path.name}", category=get("category"), website=get("website"),
                                        email=get("email"), phone=get("phone"), instagram=ig, city=get("city") or "Berlin")
        stats["leads_new" if created else "leads_updated"] += 1
        note = " · ".join(x for x in (get("contact_name") and f"Contact: {get('contact_name')}", get("notes")) if x)
        status_raw = get("status")
        lead_status, partner_stage = next(((st, ps) for rx, st, ps in STATUS_MAP if status_raw and re.search(rx, status_raw, re.I)),
                                          (None, None))
        current = db.one("SELECT status FROM leads WHERE id=?", (lid,))["status"] or "new"
        if lead_status and _moves_forward(current, lead_status):
            leads_mod.set_status(db, lid, lead_status, note)
        elif note:
            db.x("UPDATE leads SET notes=trim(coalesce(notes,'') || ' ' || ?) WHERE id=? AND coalesce(notes,'') NOT LIKE ?",
                 (note, lid, f"%{note}%"))
        if lead_status == "partner" and db.one("SELECT status FROM leads WHERE id=?", (lid,))["status"] == "partner":
            pid = ensure_partner_from_lead(db, lid)
            fields = {"contact_name": get("contact_name"), "notes": get("notes")}
            if partner_stage == "live":
                fields["stage"] = "live"
            update_partner(db, pid, **{k: v for k, v in fields.items() if v})
            stats["partners"] += 1
    return stats


def _moves_forward(current: str, new: str) -> bool:
    """An old spreadsheet must not move a lead back ("contacted" in the sheet, "meeting" in Setz)."""
    if new == "lost":
        return current not in {"partner", "lost"}
    if current == "lost":
        return new == "partner"
    rank = {s: i for i, s in enumerate(STATUS_ORDER)}
    return rank.get(new, 0) > rank.get(current, 0)


def _name_from_email(email: str, contact: str = "") -> str:
    """A company domain names the business; a freemail address (anna@gmail.com) doesn't, so use the person."""
    from .leadintel import FREEMAIL

    if "@" not in email:
        return contact
    local, dom = email.lower().rsplit("@", 1)
    if dom not in FREEMAIL:
        return dom.split(".")[0].title()
    return contact or re.sub(r"[._+-]+", " ", local).strip().title()


def _read_text(path: Path) -> str:
    if path.suffix.lower() == ".docx":
        try:
            import docx  # python-docx
        except ImportError as exc:
            raise RuntimeError("Install python-docx to import .docx (pip install python-docx), or save as .txt/.md.") from exc
        return "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
    return path.read_text(encoding="utf-8", errors="replace")


def classify_doc(name: str, text: str) -> str:
    hay = (name + " " + text[:3000]).lower()
    if re.search(r"caption|hashtag|#\w+|reel|tiktok|instagram|story|carousel|hook|content calendar|post idea", hay):
        return "social"
    if re.search(r"subject:|betreff:|dear |hi \w+,|hallo \w+|sehr geehrte|outreach|follow[- ]up|partnership|kooperation|pitch|template|cold email", hay):
        return "outreach"
    return "notes"


def import_doc(cfg: Config, path: Path) -> str | None:
    text = _read_text(path).strip()
    if len(text) < 40:
        return None
    kind = classify_doc(path.name, text)
    digest = hashlib.sha1(text.encode()).hexdigest()[:10]
    out_dir = cfg.data_dir / "library" / kind
    out_dir.mkdir(parents=True, exist_ok=True)
    if any(p.name.endswith(f"-{digest}.md") for p in out_dir.glob("*.md")):
        return None
    slug = re.sub(r"[^a-z0-9]+", "-", path.stem.lower()).strip("-")[:50] or "doc"
    (out_dir / f"{slug}-{digest}.md").write_text(f"<!-- imported from {path.name} -->\n{text}\n", encoding="utf-8")
    return kind


def run(cfg: Config, db: DB, folder: str | Path) -> dict[str, Any]:
    root = Path(folder).expanduser()
    if not root.exists():
        raise RuntimeError(f"{root} not found")
    files = [root] if root.is_file() else sorted(p for p in root.rglob("*") if p.is_file())
    report: dict[str, Any] = {"tables": {}, "docs": {"outreach": 0, "social": 0, "notes": 0}, "skipped": [], "errors": {}}
    for p in files:
        suf = p.suffix.lower()
        try:
            if suf in {".csv", ".xlsx", ".json"}:
                report["tables"][p.name] = import_table(db, p)
            elif suf in {".md", ".txt", ".docx"}:
                kind = import_doc(cfg, p)
                if kind:
                    report["docs"][kind] += 1
            else:
                report["skipped"].append(p.name)
        except Exception as exc:  # one bad file shouldn't stop the import
            report["errors"][p.name] = str(exc)[:200]
    return report


def library(cfg: Config, kind: str, max_chars: int = 6000) -> str:
    """Concatenated imported documents of one kind (newest first), for prompts."""
    d = cfg.data_dir / "library" / kind
    if not d.exists():
        return ""
    out, total = [], 0
    for p in sorted(d.glob("*.md"), key=lambda x: x.stat().st_mtime, reverse=True):
        t = p.read_text(encoding="utf-8")
        if total + len(t) > max_chars:
            t = t[: max(0, max_chars - total)]
        out.append(t)
        total += len(t)
        if total >= max_chars:
            break
    return "\n\n---\n\n".join(out)
