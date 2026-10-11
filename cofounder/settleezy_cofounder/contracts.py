"""Lex, the contracts agent: prepares partnership contracts from what Setz already knows, and reviews drafts.

Prepare
  `prepare(cfg, db, "partner_de", partner_id=7, fields={...})` fills a template with everything Setz knows
  (Impressum legal name, owner, address, the agreed offer, today's date, your company details from
  [contracts] in config.toml) and lists what is still missing. Output: Markdown + .docx in data/contracts/.
  Your own Word templates work too: put `my_template.docx` with {{ placeholders }} into data/contracts/templates/.
  When a partner reaches "Agreed in principle", Lex drafts the agreement automatically (autopilot).

Review
  `analyse(text)` reads a draft (.docx, .pdf, .txt, .md) and checks it clause by clause: parties, subject,
  benefit, term, termination and notice, renewal, liability, data protection (DSGVO), brand use, governing law,
  written form, severability, signatures. It flags risks (exclusivity, penalties, non-compete, liability
  exclusions that are void under § 309 Nr. 7 BGB, long notice periods, fees payable by Settleezy, foreign law,
  unfilled placeholders) and extracts terms (notice, durations, amounts, percentages, dates).
  With the Claude API configured it adds a written review with suggested changes.

Lex is a checklist and a first draft, not legal advice: have important contracts checked by a lawyer.
"""

from __future__ import annotations

import io
import json
import re
import zipfile
from datetime import date, timedelta
from pathlib import Path
from typing import Any

from .config import Config
from .db import DB, utcnow

TEMPLATE_DIR = Path(__file__).parent / "contract_templates"
TEMPLATES: dict[str, dict[str, Any]] = {
    "partner_de": {"file": "partner_de.md", "label": "Kooperationsvereinbarung (Partner, DE)", "lang": "de", "for": "partner"},
    "partner_en": {"file": "partner_en.md", "label": "Partnership agreement (EN)", "lang": "en", "for": "partner"},
    "university_en": {"file": "university_en.md", "label": "Buddy platform agreement (university, EN)", "lang": "en", "for": "university"},
    "nda_en": {"file": "nda_en.md", "label": "Mutual NDA (EN)", "lang": "en", "for": "any"},
}
DEFAULTS = {"term_months": "12", "renewal_months": "12", "notice_weeks": "4",
            "offer_conditions_de": "gültig gegen Vorlage der digitalen Mitgliedskarte; nicht mit anderen Aktionen kombinierbar",
            "offer_conditions_en": "valid on showing the digital membership card; cannot be combined with other offers",
            "student_group": "incoming international students", "services": "buddy matching, arrival guides, community events",
            "fee_terms": "As agreed in the attached offer.", "purpose": "a possible cooperation with Settleezy"}
PLACEHOLDER = re.compile(r"\{\{\s*([a-z_]+)\s*\}\}")


def _ensure(db: DB) -> None:
    db.conn.execute("CREATE TABLE IF NOT EXISTS contracts (id INTEGER PRIMARY KEY AUTOINCREMENT, template TEXT, title TEXT, "
                    "party TEXT, partner_id INTEGER, lead_id INTEGER, fields TEXT, missing TEXT, status TEXT DEFAULT 'draft', "
                    "path_md TEXT, path_docx TEXT, analysis TEXT, created_at TEXT, updated_at TEXT)")
    db.conn.commit()


def contracts_dir(cfg: Config) -> Path:
    d = cfg.data_dir / "contracts"
    (d / "templates").mkdir(parents=True, exist_ok=True)
    return d


def templates(cfg: Config) -> list[dict]:
    out = [{"key": k, **v, "source": "built-in", "fields": template_fields(k, cfg)} for k, v in TEMPLATES.items()]
    for f in sorted((contracts_dir(cfg) / "templates").glob("*.docx")):
        out.append({"key": "docx:" + f.stem, "file": str(f), "label": f"{f.stem} (your Word template)", "lang": "", "for": "any",
                    "source": "yours", "fields": template_fields("docx:" + f.stem, cfg)})
    for f in sorted((contracts_dir(cfg) / "templates").glob("*.md")):
        out.append({"key": "md:" + f.stem, "file": str(f), "label": f"{f.stem} (your template)", "lang": "", "for": "any",
                    "source": "yours", "fields": template_fields("md:" + f.stem, cfg)})
    return out


def _template_text(key: str, cfg: Config) -> str:
    if key in TEMPLATES:
        return (TEMPLATE_DIR / TEMPLATES[key]["file"]).read_text(encoding="utf-8")
    if key.startswith("md:"):
        return (contracts_dir(cfg) / "templates" / f"{key[3:]}.md").read_text(encoding="utf-8")
    if key.startswith("docx:"):
        return extract_text(contracts_dir(cfg) / "templates" / f"{key[5:]}.docx")
    raise ValueError(f"unknown template '{key}'; choose from {', '.join(t['key'] for t in templates(cfg))}")


def template_fields(key: str, cfg: Config) -> list[str]:
    try:
        return list(dict.fromkeys(PLACEHOLDER.findall(_template_text(key, cfg))))
    except (OSError, ValueError):
        return []


# -- prepare -----------------------------------------------------------------------------------------

def known_fields(cfg: Config, db: DB, *, partner_id: int | None = None, lead_id: int | None = None, lang: str = "de") -> dict[str, str]:
    """Everything Setz already knows about both parties, as template fields."""
    today = date.today()
    f: dict[str, str] = {
        "our_legal_name": cfg.get("contracts.legal_name", ""), "our_address": cfg.get("contracts.address", ""),
        "our_signatory": cfg.get("contracts.signatory", "") or cfg.me.get("name", ""),
        "date": today.strftime("%d.%m.%Y" if lang == "de" else "%d %B %Y"),
        "start_date": (today + timedelta(days=int(cfg.get("contracts.start_in_days", 7)))).strftime("%d.%m.%Y" if lang == "de" else "%d %B %Y"),
        "term_months": str(cfg.get("contracts.term_months", DEFAULTS["term_months"])),
        "renewal_months": str(cfg.get("contracts.renewal_months", DEFAULTS["renewal_months"])),
        "notice_weeks": str(cfg.get("contracts.notice_weeks", DEFAULTS["notice_weeks"])),
        "offer_conditions": DEFAULTS["offer_conditions_de" if lang == "de" else "offer_conditions_en"],
        "student_group": DEFAULTS["student_group"], "services": DEFAULTS["services"], "fee_terms": DEFAULTS["fee_terms"],
        "purpose": DEFAULTS["purpose"],
        "start_semester": ("Wintersemester " if lang == "de" else "winter semester ") + str(today.year if today.month < 10 else today.year + 1),
    }
    p = db.one("SELECT * FROM partners WHERE id=?", (partner_id,)) if partner_id else None
    if p and p["lead_id"] and not lead_id:
        lead_id = p["lead_id"]
    lead = db.one("SELECT * FROM leads WHERE id=?", (lead_id,)) if lead_id else None
    if lead is None and p is not None:
        lead = db.one("SELECT * FROM leads WHERE lower(name)=lower(?) LIMIT 1", (p["name"],))
    name = (p["name"] if p else None) or (lead["name"] if lead else "")
    if name:
        f["party_name"] = name
    if lead is not None:
        lead = dict(lead)
        f["partner_legal_name"] = lead.get("legal_name") or ""
        f["partner_signatory"] = lead.get("owner") or ""
        f["partner_address"] = lead.get("address") or ""
        f["venue_address"] = lead.get("address") or ""
    if p is not None:
        p = dict(p)
        if p.get("offer"):
            f["offer"] = p["offer"]
        if p.get("contact_name") and not f.get("partner_signatory"):
            f["partner_signatory"] = p["contact_name"]
    if not f.get("partner_legal_name") and name:
        f["partner_legal_name"] = name   # trading name until the Impressum is read
    return {k: v for k, v in f.items() if v not in (None, "")}


def render(text: str, fields: dict[str, str]) -> tuple[str, list[str]]:
    missing: list[str] = []

    def sub(m: re.Match) -> str:
        v = fields.get(m.group(1))
        if v in (None, ""):
            missing.append(m.group(1))
            return f"[{m.group(1).replace('_', ' ').upper()}]"
        return str(v)

    return PLACEHOLDER.sub(sub, text), list(dict.fromkeys(missing))


def _slug(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")[:40] or "contract"


def prepare(cfg: Config, db: DB, template: str, *, partner_id: int | None = None, lead_id: int | None = None,
            fields: dict[str, str] | None = None, contract_id: int | None = None) -> dict[str, Any]:
    """Create (or update) a contract draft. Missing information is marked [LIKE THIS] and listed."""
    _ensure(db)
    lang = TEMPLATES.get(template, {}).get("lang") or "de"
    known = known_fields(cfg, db, partner_id=partner_id, lead_id=lead_id, lang=lang)
    if contract_id:
        prev = db.one("SELECT * FROM contracts WHERE id=?", (contract_id,))
        if not prev:
            raise ValueError(f"no contract {contract_id}")
        known.update(json.loads(prev["fields"] or "{}"))
        template, partner_id, lead_id = prev["template"], prev["partner_id"], prev["lead_id"]
    known.update({k: str(v) for k, v in (fields or {}).items() if v not in (None, "")})
    body, missing = render(_template_text(template, cfg), known)
    party = known.get("party_name") or known.get("partner_legal_name") or "partner"
    title = f"{TEMPLATES.get(template, {}).get('label', template)}: {party}"
    now = utcnow()
    if contract_id:
        db.x("UPDATE contracts SET fields=?, missing=?, title=?, updated_at=? WHERE id=?",
             (json.dumps(known, ensure_ascii=False), json.dumps(missing), title, now, contract_id))
        cid = contract_id
    else:
        cid = db.x("INSERT INTO contracts(template,title,party,partner_id,lead_id,fields,missing,status,created_at,updated_at) "
                   "VALUES(?,?,?,?,?,?,?,?,?,?)", (template, title, party, partner_id, lead_id, json.dumps(known, ensure_ascii=False),
                                                 json.dumps(missing), "draft", now, now)).lastrowid
    base = contracts_dir(cfg) / f"{cid:04d}-{_slug(party)}-{template.replace(':', '-')}"
    md = base.with_suffix(".md")
    md.write_text(body, encoding="utf-8")
    docx = _write_docx(cfg, template, known, body, base.with_suffix(".docx"))
    analysis = analyse(body, lang=lang)
    db.x("UPDATE contracts SET path_md=?, path_docx=?, analysis=? WHERE id=?",
         (str(md), str(docx) if docx else "", json.dumps(analysis, ensure_ascii=False), cid))
    return {"id": cid, "title": title, "missing": missing, "path_md": str(md), "path_docx": str(docx) if docx else "",
            "text": body, "analysis": analysis}


def _write_docx(cfg: Config, template: str, fields: dict, body: str, out: Path) -> Path | None:
    """Your Word template (docxtpl) keeps its layout; built-in templates become a clean .docx (python-docx)."""
    try:
        if template.startswith("docx:"):
            from docxtpl import DocxTemplate  # type: ignore

            doc = DocxTemplate(str(contracts_dir(cfg) / "templates" / f"{template[5:]}.docx"))
            doc.render({k: v for k, v in fields.items()})
            doc.save(str(out))
            return out
        from docx import Document  # type: ignore

        doc = Document()
        for line in body.splitlines():
            s = line.rstrip()
            if s.startswith("# "):
                doc.add_heading(s[2:], level=0)
            elif s.startswith("## "):
                doc.add_heading(s[3:], level=2)
            elif s:
                para = doc.add_paragraph()
                for i, part in enumerate(re.split(r"\*\*", s)):
                    para.add_run(part).bold = bool(i % 2)
        doc.save(str(out))
        return out
    except ImportError:
        return None


def contracts(db: DB, status: str | None = None) -> list[dict]:
    _ensure(db)
    sql, params = "SELECT * FROM contracts", []
    if status:
        sql, params = sql + " WHERE status=?", [status]
    out = []
    for r in db.q(sql + " ORDER BY updated_at DESC", params):
        d = dict(r)
        d["fields"], d["missing"] = json.loads(d["fields"] or "{}"), json.loads(d["missing"] or "[]")
        a = json.loads(d.pop("analysis") or "{}")
        d["score"], d["risks"] = a.get("score"), len(a.get("risks", []))
        out.append(d)
    return out


def set_status(db: DB, cid: int, status: str) -> None:
    if status not in {"draft", "review", "sent", "signed", "cancelled"}:
        raise ValueError("status must be draft, review, sent, signed or cancelled")
    _ensure(db)
    row = db.one("SELECT partner_id FROM contracts WHERE id=?", (cid,))
    if not row:
        raise ValueError(f"no contract {cid}")
    db.x("UPDATE contracts SET status=?, updated_at=? WHERE id=?", (status, utcnow(), cid))
    if status == "signed" and row["partner_id"]:
        from .ops import update_partner

        p = db.one("SELECT stage, signed_at FROM partners WHERE id=?", (row["partner_id"],))
        if p:
            upd: dict[str, Any] = {"signed_at": p["signed_at"] or date.today().isoformat()}
            if p["stage"] == "agreed":
                upd["stage"] = "contract"
            update_partner(db, row["partner_id"], **upd)


def auto_prepare(cfg: Config, db: DB) -> list[dict]:
    """Draft an agreement for every partner that agreed in principle and has no contract yet."""
    _ensure(db)
    out = []
    tpl_default = cfg.get("contracts.default_template", "partner_de")
    for p in db.q("SELECT * FROM partners WHERE stage='agreed' AND status='onboarding' AND id NOT IN "
                  "(SELECT partner_id FROM contracts WHERE partner_id IS NOT NULL AND status != 'cancelled')"):
        tpl = "university_en" if p["kind"] == "university" else tpl_default
        out.append(prepare(cfg, db, tpl, partner_id=p["id"]))
    return out


# -- review --------------------------------------------------------------------------------------------

def extract_text(path: str | Path, data: bytes | None = None) -> str:
    """Text from .docx / .pdf / .txt / .md (path, or bytes with the file name for the type)."""
    p = Path(path)
    suffix = p.suffix.lower()
    raw = data if data is not None else p.read_bytes()
    if suffix == ".docx":
        try:
            from docx import Document  # type: ignore

            doc = Document(io.BytesIO(raw))
            parts = [para.text for para in doc.paragraphs]
            for t in doc.tables:
                for row in t.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            return "\n".join(parts)
        except ImportError:
            with zipfile.ZipFile(io.BytesIO(raw)) as z:
                xml = z.read("word/document.xml").decode("utf-8", "ignore")
            xml = re.sub(r"</w:p>", "\n", xml)
            return re.sub(r"<[^>]+>", "", xml)
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader  # type: ignore
        except ImportError as exc:
            raise RuntimeError("Reading PDFs needs pypdf: pip install -e .[contracts]") from exc
        return "\n".join((pg.extract_text() or "") for pg in PdfReader(io.BytesIO(raw)).pages)
    if suffix in (".txt", ".md", ""):
        return raw.decode("utf-8", "ignore")
    raise ValueError(f"can't read {suffix} files; use .docx, .pdf, .txt or .md")


# (key, label, pattern, severity when missing)
CLAUSES = [
    ("parties", "Parties", r"\b(zwischen|between)\b", "alert"),
    ("subject", "Subject / purpose", r"gegenstand|purpose|scope|vertragsgegenstand|leistungen", "warn"),
    ("benefit", "Member benefit / offer", r"rabatt|discount|vorteil|benefit|offer|angebot|%", "warn"),
    ("term", "Term / duration", r"laufzeit|term\b|beginnt|starts on|dauer|duration|runs for|läuft", "alert"),
    ("termination", "Termination & notice", r"kündig|terminat|notice", "alert"),
    ("liability", "Liability", r"haftung|haften|liab", "warn"),
    ("privacy", "Data protection (GDPR)", r"datenschutz|dsgvo|gdpr|personal data|personenbezogen", "warn"),
    ("brand", "Use of name & logo", r"marke|logo|brand|name und logo|trademark", "info"),
    ("law", "Governing law & jurisdiction", r"recht\b|law|gerichtsstand|jurisdiction", "warn"),
    ("form", "Written / text form for changes", r"schriftform|textform|in writing|text form|written form", "info"),
    ("severability", "Severability", r"unwirksam|salvatorisch|invalid|severab", "info"),
    ("signature", "Signatures", r"_{6,}|unterschrift|signature|signed by", "warn"),
]

# (key, pattern, severity, why, suggestion)
RISKS = [
    ("exclusivity", r"exklusiv|exclusiv|ausschließlich(?!e)|sole partner|alleinig", "alert",
     "Exclusivity limits which partners Settleezy can sign in this category or area.",
     "Remove it, or limit it to a short period, one category and a small radius."),
    ("penalty", r"vertragsstrafe|penalty|liquidated damages|strafzahlung", "alert",
     "A contractual penalty can cost more than the partnership is worth.", "Delete the penalty or cap it clearly."),
    ("noncompete", r"wettbewerbsverbot|non[- ]?compete|konkurrenzverbot", "alert",
     "A non-compete can block future partners.", "Delete it; partnerships for member discounts don't need one."),
    ("liability_excl", r"(haftung|liability)[^.]{0,80}(ausgeschlossen|excluded|keine haftung|no liability)", "alert",
     "A full liability exclusion is void in standard terms for gross negligence and injury (§ 309 Nr. 7 BGB).",
     "Use the standard limitation: unlimited for intent/gross negligence and injury, otherwise foreseeable typical damage."),
    ("unlimited_liability", r"(haftet|liable)[^.]{0,40}(unbegrenzt|unbeschränkt|unlimited|without limit)(?![^.]{0,80}(vorsatz|intent|grob|gross))", "warn",
     "Unlimited liability for all damage is unusual for a discount partnership.", "Limit liability to foreseeable typical damage."),
    ("fees", r"(provisionen|\bprovision (?:von|in höhe|beträgt)|commission|gebühr|\bfees?\b|vergütung|entgelt|zahlt settleezy|settleezy (pays|shall pay|zahlt))", "warn",
     "There is a payment term: check who pays whom and how much.", "Make sure Settleezy isn't paying for a member discount."),
    ("auto_renew", r"verlängert sich|automatically renew|auto[- ]?renew|stillschweigend", "info",
     "The contract renews automatically.", "Fine for partners; diarise the notice date."),
    ("foreign_law", r"(law of|recht (der|des)) (?!germany|deutschland|der bundesrepublik)[a-z]+", "warn",
     "Another country's law applies.", "Use German law and Berlin as place of jurisdiction."),
    ("data_transfer", r"(übermittelt|transfer|share|weitergabe)[^.]{0,60}(personal data|personenbezogene|member data|mitgliederdaten)", "warn",
     "Member data may be shared with the partner.", "Share no member data, or add a data processing agreement (Art. 28 GDPR)."),
    ("one_sided_termination", r"(nur|only) (der partner|the partner|partner) (kann|may|darf) (kündigen|terminate)", "warn",
     "Only one side may terminate.", "Give both sides the same right to terminate."),
]
_NEGATION = re.compile(r"\b(keine?n?|no|not|nicht|free of charge|kostenlos|without)\b")


def _negated(key: str, low: str, m: re.Match) -> bool:
    """'no commission or fees', 'keine Provisionen' are good news, not a fee clause."""
    if key not in {"fees", "data_transfer", "exclusivity", "penalty"}:
        return False
    return bool(_NEGATION.search(low[max(0, m.start() - 40):m.start()]))


PLACEHOLDERS = re.compile(r"\{\{[^}]*\}\}|\[[A-Z][A-Z _]{2,}\]|\bTBD\b|\bXX+\b|\.\.\.\s*€|__+\s*€")
DURATION = re.compile(r"(\d{1,3})\s*(wochen|woche|weeks?|monate[n]?|monat|months?|jahre[n]?|jahr|years?|tage[n]?|days?)", re.I)


def _months(n: int, unit: str) -> float:
    u = unit.lower()
    return n / 4.345 if u.startswith(("woch", "week")) else n * 12 if u.startswith(("jahr", "year")) else n / 30 if u.startswith(("tag", "day")) else n


def analyse(text: str, *, lang: str | None = None) -> dict[str, Any]:
    """Clause checklist, risks, extracted terms and a 0-100 score (higher = safer, more complete)."""
    t = text or ""
    low = t.lower()
    lang = lang or ("de" if len(re.findall(r"\b(und|der|die|das|vertrag|wird)\b", low)) > len(re.findall(r"\b(and|the|agreement|shall)\b", low)) else "en")
    clauses = [{"key": k, "label": label, "present": bool(re.search(rx, low)), "severity": sev} for k, label, rx, sev in CLAUSES]
    risks = []
    for key, rx, sev, why, fix in RISKS:
        m = next((m for m in re.finditer(rx, low) if not _negated(key, low, m)), None)
        if m:
            s, e = max(0, m.start() - 80), min(len(t), m.end() + 80)
            risks.append({"key": key, "severity": sev, "why": why, "suggestion": fix, "excerpt": re.sub(r"\s+", " ", t[s:e]).strip()})
    holes = sorted(set(PLACEHOLDERS.findall(t)))
    if holes:
        risks.append({"key": "placeholders", "severity": "alert", "why": f"{len(holes)} gaps still to fill: {', '.join(holes[:6])}",
                      "suggestion": "Fill in the missing details before sending.", "excerpt": ""})
    # notice period
    notice = None
    for m in re.finditer(r"(kündig|notice|frist)[^.]{0,120}", low):
        d = DURATION.search(m.group(0))
        if d:
            notice = {"value": int(d.group(1)), "unit": d.group(2), "months": round(_months(int(d.group(1)), d.group(2)), 1)}
            break
    if notice and notice["months"] > 3:
        risks.append({"key": "long_notice", "severity": "warn", "why": f"Notice period of {notice['value']} {notice['unit']} is long.",
                      "suggestion": "Ask for 4 weeks to 3 months.", "excerpt": ""})
    term = None
    for m in re.finditer(r"(laufzeit|läuft|runs for|term of|duration|dauer)[^.]{0,80}", low):
        d = DURATION.search(m.group(0))
        if d:
            term = {"value": int(d.group(1)), "unit": d.group(2), "months": round(_months(int(d.group(1)), d.group(2)), 1)}
            break
    if term and term["months"] > 24:
        risks.append({"key": "long_term", "severity": "warn", "why": f"Fixed term of {term['value']} {term['unit']}.",
                      "suggestion": "Start with 6-12 months and renew.", "excerpt": ""})
    terms = {
        "notice": notice, "term": term,
        "amounts": sorted(set(re.findall(r"\d[\d.,]*\s?(?:€|eur\b|euro)|€\s?\d[\d.,]*", t, re.I)))[:10],
        "percentages": sorted(set(re.findall(r"\d{1,2}(?:[.,]\d)?\s?%", t)))[:10],
        "dates": sorted(set(re.findall(r"\b\d{1,2}\.\d{1,2}\.\d{2,4}\b|\b\d{4}-\d{2}-\d{2}\b", t)))[:10],
        "statutes": sorted(set(re.findall(r"§\s?\d+[a-z]?(?:\s?(?:Abs\.|Nr\.)\s?\d+)*\s?[A-Z][A-Za-z]{1,6}", t)))[:10],
    }
    penalty = {"alert": 18, "warn": 8, "info": 2}
    score = 100
    for c in clauses:
        if not c["present"]:
            score -= penalty[c["severity"]]
    for r in risks:
        score -= penalty[r["severity"]] if r["key"] != "auto_renew" else 0
    missing = [c["label"] for c in clauses if not c["present"] and c["severity"] in ("alert", "warn")]
    verdict = ("Ready to send after a final read." if score >= 85 and not any(r["severity"] == "alert" for r in risks)
               else "Fix the flagged points before sending." if score >= 60 else "Not ready: important clauses are missing or risky.")
    return {"lang": lang, "score": max(0, score), "verdict": verdict, "clauses": clauses, "missing": missing,
            "risks": sorted(risks, key=lambda r: {"alert": 0, "warn": 1, "info": 2}[r["severity"]]), "terms": terms,
            "words": len(t.split()), "disclaimer": "Checklist and first draft, not legal advice."}


def ai_review(cfg: Config, text: str, analysis: dict | None = None) -> str:
    """Optional written review with suggested wording (Claude API). Empty string if not configured."""
    from .llm import LLM

    llm = LLM(cfg)
    if not llm.cloud_enabled or not cfg.get("contracts.ai_review", True):
        return ""
    a = analysis or analyse(text)
    try:
        return llm.cloud(
            "Contract draft (a partnership between Settleezy, a Berlin student discount membership, and a venue/brand/university):\n\n"
            + text[:24000] + "\n\nAutomatic checks found: " + json.dumps({"missing": a["missing"], "risks": [r["why"] for r in a["risks"]]})
            + "\n\nReview it for Settleezy: 1) the 3-5 most important issues, each with a concrete replacement wording "
              "(same language as the contract), 2) anything unusual for German contract law (AGB, §§ 305-310 BGB), "
              "3) a one-line verdict. Be concise. Remind that this is not legal advice in one short line.",
            "You are Lex, a careful contracts assistant for a German start-up.", effort="medium", max_tokens=3000)
    except Exception as exc:
        return f"(AI review unavailable: {exc})"


def review_file(cfg: Config, db: DB, filename: str, data: bytes | None = None, *, ai: bool = False,
                partner_id: int | None = None) -> dict[str, Any]:
    """Analyse an uploaded / local draft and keep it in the contracts list (status 'review')."""
    _ensure(db)
    text = extract_text(filename, data)
    a = analyse(text)
    if ai:
        a["ai"] = ai_review(cfg, text, a)
    dest = contracts_dir(cfg) / "reviews"
    dest.mkdir(exist_ok=True)
    name = Path(filename).name
    if data is not None:
        (dest / name).write_bytes(data)
    now = utcnow()
    cid = db.x("INSERT INTO contracts(template,title,party,partner_id,fields,missing,status,path_md,analysis,created_at,updated_at) "
               "VALUES(?,?,?,?,?,?,?,?,?,?,?)", ("review", f"Review: {name}", name, partner_id, "{}", json.dumps(a["missing"]), "review",
                                                str(dest / name) if data is not None else str(filename),
                                                json.dumps(a, ensure_ascii=False), now, now)).lastrowid
    return {"id": cid, **a}


def get(db: DB, cid: int) -> dict[str, Any]:
    _ensure(db)
    r = db.one("SELECT * FROM contracts WHERE id=?", (cid,))
    if not r:
        raise ValueError(f"no contract {cid}")
    d = dict(r)
    d["fields"], d["missing"], d["analysis"] = json.loads(d["fields"] or "{}"), json.loads(d["missing"] or "[]"), json.loads(d["analysis"] or "{}")
    try:
        d["text"] = Path(d["path_md"]).read_text(encoding="utf-8") if d["path_md"] and d["path_md"].endswith(".md") else ""
    except OSError:
        d["text"] = ""
    return d
